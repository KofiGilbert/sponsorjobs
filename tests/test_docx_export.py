"""Tests for the ATS-friendly .docx export (tailoring/docx_export.py).

The .docx must be a real, openable Word file built from the SAME profile the PDF uses,
and it must be PARSER-SAFE: single column, no tables, no images, contact in the body.
"""

from __future__ import annotations

import pytest

docx = pytest.importorskip("docx")   # python-docx is a declared dep; skip if absent
from docx import Document

from tailoring.docx_export import build_docx, profile_to_plaintext

PROFILE = {
    "identity": {"name": "Ada Lovelace", "email": "ada@example.com",
                 "phone": "555-0100", "linkedin": "https://linkedin.com/in/ada",
                 "github": "https://github.com/ada"},
    "summary": "Backend engineer with distributed-systems experience.",
    "education": [{"school": "Trinity College", "location": "London",
                   "degree": "BSc Computer Science", "date": "2024",
                   "courses": "Algorithms, Databases"}],
    "skills": {"Languages": "Python, Go", "Cloud": "AWS, Docker"},
    "experience": [{"org": "Acme Corp", "location": "New York, NY",
                    "roles": [{"title": "SWE Intern", "dates": "Summer 2023",
                               "bullets": ["Built an ETL pipeline processing 2M rows/day",
                                           "Cut API latency 40% via caching"]}]}],
    "projects": [{"org": "PocketLedger", "link": "https://github.com/ada/pocketledger",
                  "bullets": ["A local-first budgeting app in Go"]}],
}


def test_build_docx_is_a_valid_openable_word_file(tmp_path):
    out = build_docx(PROFILE, tmp_path / "cv.docx")
    doc = Document(out)                      # re-opens without error => valid .docx
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Ada Lovelace" in text
    assert "EXPERIENCE" in text and "SKILLS" in text
    assert "Built an ETL pipeline processing 2M rows/day" in text
    assert "ada@example.com" in text        # contact is in the BODY, not a header


def test_docx_is_ats_parser_safe(tmp_path):
    """No tables, no images, no header/footer content — the layout choices that scramble
    ATS parsing. Everything the person needs read is in the single-column body."""
    out = build_docx(PROFILE, tmp_path / "cv.docx")
    doc = Document(out)
    assert doc.tables == [], "an ATS-friendly resume must have no tables"
    # No inline images.
    assert not doc.inline_shapes, "no images in an ATS resume"
    # Headers/footers carry no text (parsers skip them, so nothing important may live there).
    for section in doc.sections:
        assert "".join(p.text for p in section.header.paragraphs).strip() == ""
        assert "".join(p.text for p in section.footer.paragraphs).strip() == ""


def test_bullets_are_real_list_paragraphs(tmp_path):
    out = build_docx(PROFILE, tmp_path / "cv.docx")
    doc = Document(out)
    bulleted = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert any("ETL pipeline" in p.text for p in bulleted)


def test_plaintext_preview_keeps_content_in_order():
    txt = profile_to_plaintext(PROFILE)
    assert txt.startswith("Ada Lovelace")
    # Section order: summary before experience before the bullet content.
    assert txt.index("SUMMARY") < txt.index("EXPERIENCE")
    assert "• Built an ETL pipeline processing 2M rows/day" in txt
    assert "Languages: Python, Go" in txt


def test_empty_profile_does_not_crash(tmp_path):
    out = build_docx({}, tmp_path / "empty.docx")
    Document(out)                             # still a valid, if sparse, file
    assert profile_to_plaintext({}) is not None


def test_flat_role_entry_without_roles_list(tmp_path):
    """An entry may carry title/dates/bullets directly instead of a `roles` list."""
    prof = {"identity": {"name": "X"},
            "experience": [{"org": "Solo", "title": "Founder", "dates": "2022-2024",
                            "bullets": ["Shipped a product"]}]}
    out = build_docx(prof, tmp_path / "flat.docx")
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Founder" in text and "Shipped a product" in text


def test_build_cover_docx_has_header_and_body(tmp_path):
    from tailoring.docx_export import build_cover_docx
    ident = {"name": "Ada Lovelace", "email": "ada@example.com"}
    body = "Dear Hiring Team,\n\nI am excited to apply.\n\nSincerely,\nAda"
    out = build_cover_docx(ident, body, tmp_path / "cover.docx")
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Ada Lovelace" in text and "ada@example.com" in text     # contact header in body
    assert doc.tables == []                                         # ATS-safe: no tables
    # Blank-line-separated blocks become separate paragraphs.
    assert any(p.text.strip() == "I am excited to apply." for p in doc.paragraphs)
