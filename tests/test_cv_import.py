"""Acceptance tests for CV upload + extraction + OCR + the required-field gate
(CLAUDE.md §4a, §8). Everything is local and offline.

Covers:
  * text extraction from a generated .docx and a text PDF,
  * OCR extraction from a generated image (gated: requires RapidOCR),
  * the clean tool-authored summary is written to memory while the RAW extracted
    text is NOT indexed,
  * the raw uploaded file is kept on disk as a backup and not deleted,
  * the flow refuses to finalize when a required field (a job's dates) is missing,
  * bullets are drafted for a real role given only its title + company.
"""

from __future__ import annotations

from pathlib import Path

from intake.cv_import import extract_profile, extract_text, summarize_import
from intake.memory import ConversationMemory
from intake.palace_memory import PalaceMemory
from ui.records import CVRecords
from ui.session import WebIntake
from conftest import requires_latex, requires_ocr

JD = "Data Analyst - Globex - Chicago. SQL, Python, dashboards, statistics."

# A CV body phrased so the offline FakeLLM extractor can parse it, while reading
# like real CV content.
CV_BODY = (
    "I'm Jane Doe, jane@example.com, https://github.com/janedoe .\n"
    "I did an M.S. in Statistics at Northwestern in 2017.\n"
    "I work as a Data Analyst at Globex since Jan 2019.\n"
)


def _mk(template_source, fake_llm, workdir, palace_dir, jd=JD):
    return WebIntake(jd, template_source, fake_llm, ConversationMemory(":memory:"),
                     CVRecords(":memory:"), workdir, jobname="cv", palace_dir=palace_dir)


# -- extraction --------------------------------------------------------- #

def test_extract_from_docx(fake_llm, tmp_path):
    import docx
    p = tmp_path / "cv.docx"
    doc = docx.Document()
    for line in CV_BODY.strip().splitlines():
        doc.add_paragraph(line)
    doc.save(str(p))

    text = extract_text(p)
    assert "Jane Doe" in text and "Globex" in text and "Northwestern" in text

    ess = extract_profile(p, fake_llm)
    assert ess["identity"]["name"] == "Jane Doe"
    assert ess["experience"][0]["org"] == "Globex"
    assert ess["experience"][0]["title"] == "Data Analyst"
    assert ess["education"][0]["date"] == "2017"


def test_extract_from_text_pdf(fake_llm, tmp_path):
    import pytest
    reportlab = pytest.importorskip("reportlab.pdfgen.canvas")
    p = tmp_path / "cv.pdf"
    c = reportlab.Canvas(str(p))
    y = 760
    for line in CV_BODY.strip().splitlines():
        c.drawString(72, y, line)
        y -= 24
    c.showPage()
    c.save()

    text = extract_text(p)
    assert "Jane Doe" in text and "Globex" in text

    ess = extract_profile(p, fake_llm)
    assert ess["identity"]["name"] == "Jane Doe"
    assert ess["experience"][0]["org"] == "Globex"


@requires_ocr
def test_ocr_from_image(fake_llm, tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1000, 340), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 46)
    except Exception:
        font = ImageFont.load_default()
    d.text((30, 30), "Jane Doe", fill="black", font=font)
    d.text((30, 120), "Data Analyst at Globex", fill="black", font=font)
    d.text((30, 210), "MBA from DePaul 2019", fill="black", font=font)
    p = tmp_path / "scan.png"
    img.save(str(p))

    text = extract_text(p)
    low = text.lower()
    assert "jane" in low and "globex" in low   # a scanned/image CV is readable


# -- storage rules ------------------------------------------------------ #

def test_clean_summary_indexed_but_raw_text_not(template_source, fake_llm, tmp_path):
    palace = tmp_path / "palace"
    s = _mk(template_source, fake_llm, tmp_path / "b", palace)
    s.start()
    raw = CV_BODY + "\nZZRAWONLYTOKEN garbled ocr noise line\n"
    s.ingest_resume(raw, "resume.docx")

    mem_text = " ".join(t["content"] for t in s.palace.load_history())
    assert "uploaded a resume" in mem_text and "Globex" in mem_text  # clean summary indexed
    assert "ZZRAWONLYTOKEN" not in mem_text                          # raw text NOT indexed
    # and the summary is a real, tool-authored sentence
    summ = summarize_import(s.essentials, "resume.docx", "2026-07-11")
    assert summ.startswith("User uploaded a resume") and "Globex" in summ


@requires_latex
def test_raw_upload_file_kept_and_not_deleted(fake_llm, tmp_path, monkeypatch):
    import io

    import docx

    import ui.app as app
    monkeypatch.setattr(app, "DB_PATH", str(tmp_path / "db.sqlite"))
    monkeypatch.setattr(app, "WORKDIR", tmp_path / "build")
    monkeypatch.setattr(app, "PALACE_DIR", tmp_path / "palace")
    monkeypatch.setattr(app, "UPLOADS_DIR", tmp_path / "uploads")
    monkeypatch.setenv("RESUME_AGENT_FAKE_LLM", "1")

    client = app.app.test_client()
    assert client.post("/api/session/start", json={"jd": JD}).status_code == 200

    buf = io.BytesIO()
    doc = docx.Document()
    for line in CV_BODY.strip().splitlines():
        doc.add_paragraph(line)
    doc.save(buf)
    buf.seek(0)

    r = client.post("/api/session/upload_cv",
                    data={"file": (buf, "myresume.docx")},
                    content_type="multipart/form-data")
    assert r.status_code == 200

    kept = list((tmp_path / "uploads").glob("*myresume.docx"))
    assert len(kept) == 1 and kept[0].exists()          # raw backup written
    # a second, unrelated action must not remove it
    client.post("/api/session/answer", json={"text": "thanks"})
    assert kept[0].exists()                              # never deleted


# -- the gate ----------------------------------------------------------- #

@requires_latex
def test_gate_refuses_to_finalize_without_job_dates(template_source, fake_llm, workdir):
    s = _mk(template_source, fake_llm, workdir, workdir / "palace")
    s.start()
    # name + email + dated degree, but the job has NO dates
    st = s.submit("I'm Jane Doe, jane@example.com. M.S. in Statistics at Northwestern "
                  "in 2017. I work as a Data Analyst at Globex.")
    assert st["phase"] == "chatting"          # did NOT finalize
    assert s.assembled is None                # no CV built
    joined = " ".join(st["messages"]).lower()
    assert "dates" in joined and "globex" in joined   # asks for the missing dates, by name

    # The gate holds while the person keeps answering normally — it only relaxes on
    # an EXPLICIT opt-out (tested in test_template_fidelity).
    st2 = s.submit("what about my summary section")   # not an opt-out
    assert st2["phase"] == "chatting"
    assert s.assembled is None


def test_bullets_drafted_for_role_given_title_and_company(fake_llm):
    essentials = {"identity": {"name": "Jane Doe"}, "education": [],
                  "experience": [{"org": "Globex", "title": "Data Analyst", "dates": ""}]}
    profile = fake_llm.draft_profile(JD, essentials)
    role = profile["experience"][0]
    assert role["org"] == "Globex"                         # real role kept as given
    assert role["bullets"] and all(role["bullets"])        # descriptive bullets drafted
